# 封面格式修复：把"空格+字符下划线"的字段区（段落9-17）替换为无边框表格，
# 值单元格只保留下框线 —— 线粗/线长统一、文字居中，PDF 转换稳定。
# 字段内容与原封面完全一致（从原段落提取核对过）。
import shutil, time
from docx import Document
from docx.oxml import parse_xml

PATH = 'docs/基于Mamba状态空间模型的股票价格预测方法研究-中期.docx'
shutil.copy(PATH, 'docs/中期_v18_backup_' + time.strftime('%Y%m%d_%H%M%S') + '.bak')
d = Document(PATH)

# 页面内容宽度（twips）
sec = d.sections[0]
content_w = int((sec.page_width - sec.left_margin - sec.right_margin) / 635)
# 列宽分配：标签1 / 值1 / 标签2 / 值2
lw1, vw1, lw2 = 1700, 2600, 1250
vw2 = content_w - lw1 - vw1 - lw2

FONT = '<w:rFonts w:ascii="SimSun" w:hAnsi="SimSun" w:eastAsia="宋体"/><w:b/><w:bCs/><w:sz w:val="30"/><w:szCs w:val="30"/>'

def cell(text, width, underline, align='center', span=1):
    border = ('<w:tcBorders><w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/></w:tcBorders>'
              if underline else '<w:tcBorders><w:bottom w:val="nil"/></w:tcBorders>')
    span_xml = f'<w:gridSpan w:val="{span}"/>' if span > 1 else ''
    return (f'<w:tc><w:tcPr><w:tcW w:w="{width}" w:type="dxa"/>{span_xml}'
            f'{border}<w:vAlign w:val="bottom"/></w:tcPr>'
            f'<w:p><w:pPr><w:spacing w:line="300" w:lineRule="auto"/><w:jc w:val="{align}"/>'
            f'<w:rPr>{FONT}</w:rPr></w:pPr>'
            f'<w:r><w:rPr>{FONT}</w:rPr><w:t xml:space="preserve">{text}</w:t></w:r></w:p></w:tc>')

def row(cells_xml):
    return (f'<w:tr><w:trPr><w:trHeight w:val="560" w:hRule="atLeast"/></w:trPr>{cells_xml}</w:tr>')

W3 = vw1 + lw2 + vw2  # 单字段行的合并值宽度
rows = [
    # 报告题目（两行，与原版式一致）
    row(cell('报告题目', lw1, False, 'left') + cell('基于Mamba状态空间模型的', W3, True, span=3)),
    row(cell('', lw1, False) + cell('股票价格预测方法研究', W3, True, span=3)),
    # 双字段行
    row(cell('学生姓名', lw1, False, 'left') + cell('***', vw1, True)
        + cell('学号', lw2, False) + cell('***', vw2, True)),
    row(cell('指导教师', lw1, False, 'left') + cell('***', vw1, True)
        + cell('职称', lw2, False) + cell('副教授', vw2, True)),
    # 单字段行
    row(cell('学位类别', lw1, False, 'left') + cell('电子信息硕士', W3, True, span=3)),
    row(cell('学科专业', lw1, False, 'left') + cell('人工智能', W3, True, span=3)),
    row(cell('研究方向', lw1, False, 'left') + cell('模式识别与机器学习', W3, True, span=3)),
    row(cell('研究所（院系）', lw1, False, 'left') + cell('人工智能学院', W3, True, span=3)),
    row(cell('填表日期', lw1, False, 'left') + cell('2026.10.07', W3, True, span=3)),
]

tbl_xml = (
    '<w:tbl xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    f'<w:tblPr><w:tblW w:w="{content_w}" w:type="dxa"/><w:jc w:val="center"/>'
    '<w:tblLayout w:type="fixed"/>'
    '<w:tblBorders>'
    '<w:top w:val="none" w:sz="0" w:space="0"/><w:left w:val="none" w:sz="0" w:space="0"/>'
    '<w:bottom w:val="none" w:sz="0" w:space="0"/><w:right w:val="none" w:sz="0" w:space="0"/>'
    '<w:insideH w:val="none" w:sz="0" w:space="0"/><w:insideV w:val="none" w:sz="0" w:space="0"/>'
    '</w:tblBorders>'
    '<w:tblCellMar><w:left w:w="108" w:type="dxa"/><w:right w:w="108" w:type="dxa"/></w:tblCellMar>'
    '</w:tblPr>'
    f'<w:tblGrid><w:gridCol w:w="{lw1}"/><w:gridCol w:w="{vw1}"/><w:gridCol w:w="{lw2}"/><w:gridCol w:w="{vw2}"/></w:tblGrid>'
    + ''.join(rows) + '</w:tbl>'
)

# 先收集待删段落元素（段落9-17 = 9个字段段），再把表格插到段落8之后
victims = [d.paragraphs[i]._p for i in range(9, 18)]
d.paragraphs[8]._p.addnext(parse_xml(tbl_xml))
for el in victims:
    el.getparent().remove(el)

d.save(PATH)
print('封面表格化完成，内容宽度', content_w, 'twips')
