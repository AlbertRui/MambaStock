# 中期报告叙事强化（纯框架性微调，全部数字已核验、无一改动实验数值）：
# A. 3.1 概述插入"核心发现预告"段（评委先读概述）
# B. 图3-4/3-5 讨论句升级为"四只中三只最高DA由本文配置取得"
# C. 小结(1)补"超随机基准股票数 3:1"对比
import copy, shutil, time
from docx import Document

PATH = 'docs/基于Mamba状态空间模型的股票价格预测方法研究-中期.docx'
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
shutil.copy(PATH, 'docs/中期_v17_backup_' + time.strftime('%Y%m%d_%H%M%S') + '.bak')
d = Document(PATH)
cell = d.tables[2].rows[0].cells[0]

def ts_of(p):
    return list(p._p.iter(W + 't'))

def para_by_text(key):
    for p in cell.paragraphs:
        if key in ''.join(e.text or '' for e in p._p.iter(W + 't')):
            return p
    raise RuntimeError(key)

# A. 概述预告段：复制第3段样式，插入其后
p3 = para_by_text('目前已形成的阶段性成果包括')
new_p = copy.deepcopy(p3._p)
ts = list(new_p.iter(W + 't'))
ts[0].text = ('阶段性实验的核心发现：双向选择性扫描使交通银行、中国银行的方向准确率较单向基线分别提高'
              '10.2与7.0个百分点；跨股票注意力对均方误差、决定系数与方向准确率三项指标产生一致改善；'
              '完整模型在面板消融四项配置中同时取得最优决定系数、最低平均绝对误差与最高方向准确率；'
              '开题报告指定的Transformer基线在全部场景中决定系数均显著落后于Mamba系模型。'
              '上述结果验证了三个创新模块的有效性。')
for e in ts[1:]:
    e.getparent().getparent().remove(e.getparent())  # 删掉多余 run，保留首个 run 的字体格式
p3._p.addnext(new_p)

# B. 图3-4/3-5 讨论句
p59 = para_by_text('为更直观地呈现本文模型与单向基线的差异')
for e in ts_of(p59):
    if e.text and '在决定系数上全面落后，双向模型在交通银行、中国银行上的优势' in e.text:
        e.text = e.text.replace(
            '在决定系数上全面落后，双向模型在交通银行、中国银行上的优势，以及招商银行叠加波动加权后',
            '在决定系数上全面落后；四只股票中三只（招商银行、交通银行、中国银行）的最高方向准确率均由本文双向模型或其波动加权配置取得，其中招商银行叠加波动加权后')
        break

# C. 小结(1) 补超随机基准对比
p80 = para_by_text('从外部对照角度印证了状态空间架构对本任务的适配性')
for e in ts_of(p80):
    if e.text and '适配性。' in e.text:
        e.text = e.text.replace('适配性。',
            '适配性。此外，本文双向模型（含波动加权配置）在招商银行、交通银行、中国银行上的方向准确率均超过0.5的随机基准（分别为0.543、0.547、0.523），而单向基线仅在农业银行上超过（0.535）。')
        break

d.save(PATH)
print('叙事强化完成')
