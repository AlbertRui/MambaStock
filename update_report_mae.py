# 中期报告 MAE 增量更新（v2 口径数据，数值不动）：
# 1) 表3-3/表3-4 增加 MAE(×10⁻³) 列  2) 新增图3-6（MSE+MAE 双子图），原图3-6~3-8 顺移
# 3) 正文图3-4/3-5 讨论段补 MAE 对比句（结论经真实数据核验）
import copy, re, shutil, time
from docx import Document

PATH = 'docs/基于Mamba状态空间模型的股票价格预测方法研究-中期.docx'
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
W14 = '{http://schemas.microsoft.com/office/word/2010/wordml}'
R = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
V = '{urn:schemas-microsoft-com:vml}'

shutil.copy(PATH, 'docs/中期_v15_backup_' + time.strftime('%Y%m%d_%H%M%S') + '.bak')
d = Document(PATH)
cell = d.tables[2].rows[0].cells[0]

def para_text(p):
    return ''.join(e.text or '' for e in p.iter(W + 't'))

# ---------- 1) 图号顺移（降序，仅当同段前一个 w:t 以"图"结尾） ----------
for old, new in (('3-8', '3-9'), ('3-7', '3-8'), ('3-6', '3-7')):
    for p in cell.paragraphs:
        ts = list(p._p.iter(W + 't'))
        for j, e in enumerate(ts):
            txt = e.text or ''
            if j > 0 and (ts[j-1].text or '').endswith('图'):
                if txt == old:
                    e.text = new
                elif txt == old + ' ':
                    e.text = new + ' '

# ---------- 2) 正文补 MAE 句 ----------
NEW_SENT = ('图3-6进一步给出均方误差与平均绝对误差的分组对比：'
            '各股票的最优配置在两项误差指标下保持一致，且与决定系数的排序格局相同。')
for p in cell.paragraphs:
    if '图3-4与图3-5给出单股票场景' in para_text(p._p):
        for e in p._p.iter(W + 't'):
            if e.text and e.text.endswith('个百分点。'):
                e.text = e.text + NEW_SENT
                break
        break

# ---------- 3) 表格加 MAE 列 ----------
def set_cell_text(tc, text):
    ts = list(tc.iter(W + 't'))
    ts[0].text = text
    for e in ts[1:]:
        e.getparent().remove(e)

def add_col_after_mse(tbl, value_of):
    # tblGrid 同步加列
    grid = tbl._tbl.find(W + 'tblGrid')
    cols = grid.findall(W + 'gridCol')
    grid.insert(list(grid).index(cols[2]) + 1, copy.deepcopy(cols[2]))
    hdr_done = False
    for row in tbl.rows:
        tcs = row._tr.findall(W + 'tc')
        new_tc = copy.deepcopy(tcs[2])
        tcs[2].addnext(new_tc)
        if not hdr_done:
            set_cell_text(new_tc, 'MAE(×10⁻³)')
            hdr_done = True
        else:
            set_cell_text(new_tc, value_of(row))

MAE33 = {
    ('600036', '单向基线'): '9.11', ('600036', 'Transformer'): '9.80',
    ('600036', '双向'): '8.77', ('600036', '单向基线＋波动加权'): '9.10',
    ('600036', '双向＋波动加权'): '9.29',
    ('601288', '单向基线'): '10.12', ('601288', 'Transformer'): '10.92',
    ('601288', '双向'): '10.76', ('601288', '单向基线＋波动加权'): '10.18',
    ('601288', '双向＋波动加权'): '10.27',
    ('601328', '单向基线'): '9.20', ('601328', 'Transformer'): '9.74',
    ('601328', '双向'): '8.72', ('601328', '单向基线＋波动加权'): '9.21',
    ('601328', '双向＋波动加权'): '9.74',
    ('601988', '单向基线'): '8.87', ('601988', 'Transformer'): '9.56',
    ('601988', '双向'): '9.00', ('601988', '单向基线＋波动加权'): '9.52',
    ('601988', '双向＋波动加权'): '9.05',
}
def v33(row):
    cs = row.cells
    stock, cfg = cs[0].text, cs[1].text
    code = re.search(r'(\d{6})', stock).group(1)
    return MAE33[(code, cfg)]

MAE34 = {'Transformer': '9.55', '无注意力': '8.82', '＋跨股票注意力': '8.84',
         '无注意力＋波动加权': '8.85', '完整模型（注意力＋波动加权）': '8.80'}

add_col_after_mse(cell.tables[2], v33)
add_col_after_mse(cell.tables[3], lambda row: MAE34[row.cells[0].text])

# ---------- 4) 插入新图3-6 ----------
# 定位图3-5 图注段与其前面的图片段
cap35 = img35 = None
paras = cell.paragraphs
for i, p in enumerate(paras):
    if para_text(p._p).startswith('图3-5 '):
        cap35 = p._p
        img35 = paras[i - 1]._p
        break
assert cap35 is not None and img35 is not None

# 新图片段：复制图3-5 图片段，换 rId 与尺寸
rId_new, _ = d.part.get_or_add_image('results/v3_mse_mae_compare.png')
new_img = copy.deepcopy(img35)
for el in new_img.iter():
    for a in list(el.attrib):
        if a.startswith(W14):
            del el.attrib[a]
shape = new_img.find('.//' + V + 'shape')
shape.set('id', 'Picture MAE3'); shape.set('{urn:schemas-microsoft-com:office:office}spid', '_x0000_i1099')
shape.set('style', re.sub(r'height:[\d.]+pt', 'height:161.28pt', shape.get('style')))  # 403.2pt 宽 × 1680:672 比例
new_img.find('.//' + V + 'imagedata').set(R + 'id', rId_new)

# 新图注段：复制图3-5 图注，改文字
new_cap = copy.deepcopy(cap35)
for a in list(new_cap.attrib):
    if a.startswith(W14):
        del new_cap.attrib[a]
ts = list(new_cap.iter(W + 't'))
ts[1].text = '3-6 '
ts[2].text = '单股票场景四模型均方误差与平均绝对误差对比（测试集）'
for e in ts[3:]:  # 去掉原图注里的 '0.5'、'）' 等尾巴 run
    r = e.getparent()
    r.getparent().remove(r)

cap35.addnext(new_cap)
cap35.addnext(new_img)

d.save(PATH)
print('OK: MAE列 + 图3-6插入 + 图号顺移完成')
